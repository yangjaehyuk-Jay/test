# -*- coding: utf-8 -*-
"""
월별 추이 조회 (scrape.py 확장판)

scrape.py 는 '올해 1월 1일 ~ 오늘' 을 한 번만 조회한다.
이 프로그램은 시작일을 올해 1월 1일로 못박고, 종료일만 2월 1일 -> 3월 1일 ->
... -> 이번 달 1일 로 옮겨가며 같은 기업을 여러 번 조회한다.
화면의 조회 결과는 '기간 내' 기준이므로, 종료일을 늘려가며 찍으면 그 기업의
파생거래 잔액이 올해 들어 어떻게 늘고 줄었는지 추이가 나온다.

    2026-01-01 ~ 2026-02-01   (1월까지)
    2026-01-01 ~ 2026-03-01   (2월까지)
    ...
    2026-01-01 ~ 2026-09-01   (8월까지)

기본 대상은 companies2.csv 다.

로그인/조회/엑셀서식은 scrape.py 것을 그대로 가져다 쓴다. 이 파일에는
'기간을 여러 개 돌린다'는 것과 '결과를 추이표로 편다'는 것만 들어있다.

실행:  python scrape_monthly.py
       python scrape_monthly.py --companies companies.csv --anchor last
       python scrape_monthly.py --resume            (끊긴 조회 이어서)
"""
import sys, csv, time, datetime, pathlib, argparse, calendar

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

from playwright.sync_api import sync_playwright
import scrape as sc
import nexacro as nx

BASE = sc.BASE
OUT_DIR = sc.OUT_DIR
LOG_DIR = sc.LOG_DIR
DEFAULT_COMPANIES = BASE / "companies2.csv"

DELAY_SEC = 2.5

# 추이 시트 구성. (시트이름, 설명, 더할 원본 열들, 나눌 단위)
# scrape.py 의 SUMCOLS 와 같은 셈법이다. 옵션은 매수 포지션만 넣는다.
TREND_SHEETS = [
    ("매수추이", "매수 잔액 (USD mn)",
     ("선물환_매수잔액", "통화스왑_매수잔액", "옵션_매수콜"), sc.SUM_UNIT),
    ("매도추이", "매도 잔액 (USD mn)",
     ("선물환_매도잔액", "통화스왑_매도잔액", "옵션_매수풋"), sc.SUM_UNIT),
    ("건수추이", "유효 건수",
     ("선물환_건수", "옵션_건수", "통화스왑_건수"), 1),
]

# 추이표의 빈 칸이 왜 비었는지 '상태' 시트에 남긴다.
STATUS_TEXT = {sc.OK_CD: "자료있음", "401": "자료없음", sc.ERR_CD: "조회실패"}


# --------------------------------------------------------------------------
def period_ends(year, anchor="first", start_month=2, upto=None,
                include_today=False):
    """조회 종료일 목록. 지정한 달부터 이번 달까지 한 달에 하나.

    anchor='first' -> 매월 1일  (2026-02-01, 2026-03-01, ...)
    anchor='last'  -> 매월 말일 (2026-02-28, 2026-03-31, ...)
    오늘을 넘는 날짜는 넣지 않는다.
    """
    today = upto or datetime.date.today()
    out = []
    for m in range(start_month, 13):
        day = calendar.monthrange(year, m)[1] if anchor == "last" else 1
        d = datetime.date(year, m, day)
        if d > today:
            break
        out.append(d)
    if include_today and today not in out:
        out.append(today)
    return out


def ymd(d):
    return d.strftime("%Y%m%d")


def iso(d):
    return d.strftime("%Y-%m-%d")


def parse_ymd(s):
    return datetime.datetime.strptime(s, "%Y%m%d").date()


def load_done(path, log):
    """중단된 조회의 중간저장 CSV 를 읽어 (기준일, 사업자등록번호) 를 모은다."""
    done, rows = set(), []
    if not path.exists():
        return done, rows
    import io
    text, _ = sc.read_text_any(path)
    for rec in csv.DictReader(io.StringIO(text)):
        key = (rec.get("기준일"), rec.get("사업자등록번호"))
        if not key[0] or not key[1] or key in done:
            continue
        done.add(key)
        rows.append(rec)
    if rows:
        log("  이어받기: %s 에서 %d행을 되살렸습니다." % (path.name, len(rows)))
    return done, rows


# --------------------------------------------------------------------------
def run_trend(req_ctx, user, companies, st, ends, log, notify, state):
    """기간별로 전체 기업을 한 바퀴씩 돈다.

    기간을 바깥 고리에 두는 이유: 중간에 끊겨도 '어느 시점까지는 전 기업이
    다 찍혔다'가 되므로 추이표가 온전한 열 단위로 남는다.
    """
    rows, failed, done = state["rows"], state["failed"], state["done"]
    pw, part_fh = state["writer"], state["fh"]

    total = len(companies) * len(ends)
    n = 0
    sc.audit_log(req_ctx, user, "화면오픈", log)
    for ed in ends:
        log("\n  === 기준일 %s (조회기간 %s ~ %s) ==="
            % (iso(ed), iso(parse_ymd(st)), iso(ed)))
        for bizno, memo in companies:
            n += 1
            b = sc.fmt_bizno(bizno)
            if (iso(ed), b) in done:
                log("  [%3d/%d] %s - 이미 받은 값, 건너뜀" % (n, total, b))
                continue
            label = b + (" (%s)" % memo if memo else "")
            log("  [%3d/%d] %s" % (n, total, label))
            notify("%s  %d/%d  %s" % (iso(ed), n, total, b))
            try:
                data = sc.query_one(req_ctx, user, bizno, st, ymd(ed), log)
                sc.audit_log(req_ctx, user, "내역조회", log)
                rec = {"기준일": iso(ed), "사업자등록번호": b, "메모": memo or ""}
                for col, key in sc.COLMAP:
                    if key == "_bizno" or col == "사업자등록번호":
                        continue
                    rec[col] = data.get(key)
                rows.append(rec)
                log("       -> %s / %s | 선물환 %s건, 옵션 %s건, 스왑 %s건"
                    % (data.get("RSP_CD"),
                       (data.get("KIKO_CUST_NM") or "").strip(),
                       data.get("FX_VLD_CNT"), data.get("OP_VLD_CNT"),
                       data.get("CRS_VLD_CNT")))
            except Exception as e:
                log("       -> 실패: %s" % e)
                failed.append((iso(ed), b, memo or "", str(e)[:160]))
                # 실패해도 행을 남긴다. 시점 x 기업 격자가 비면 나중에
                # '조회를 안 한 것'과 '자료가 없는 것'을 구분할 수 없다.
                rec = {"기준일": iso(ed), "사업자등록번호": b, "메모": memo or "",
                       "결과코드": sc.ERR_CD, "결과내용": str(e)[:160]}
                rows.append(rec)
            pw.writerow(rec)
            part_fh.flush()
            if n < total:
                time.sleep(DELAY_SEC)
    return rows, failed


# --------------------------------------------------------------------------
def _pivot(df, companies, ends, srcs, unit):
    """기업 x 기준일 행렬. 결과코드가 000 이 아닌 칸은 비워 둔다."""
    v = df[list(srcs)].sum(axis=1, min_count=1)
    if unit != 1:
        v = v / unit
    work = df[["사업자등록번호", "기준일"]].copy()
    work["값"] = v.where(df["결과코드"].eq(sc.OK_CD))
    m = work.pivot_table(index="사업자등록번호", columns="기준일", values="값",
                         aggfunc="last")
    return m.reindex(index=[sc.fmt_bizno(b) for b, _ in companies],
                     columns=[iso(d) for d in ends])


def _name_map(df, companies):
    """사업자등록번호 -> 표시용 기업명. 응답이 준 이름이 있으면 그걸 쓴다."""
    out = {sc.fmt_bizno(b): (memo or "") for b, memo in companies}
    if "기업명" in df.columns:
        for b, nm in zip(df["사업자등록번호"], df["기업명"]):
            nm = str(nm or "").strip()
            if nm and nm not in ("nan", "None"):
                out[b] = nm
    return out


def _style_trend(ws, frame, unit):
    """추이 시트 서식. 날짜 열은 숫자, 마지막 증감 열과 총계 행은 강조."""
    from openpyxl.styles import Alignment, Font, PatternFill

    ncol, nrow = len(frame.columns), len(frame)
    last = frame.columns[-1]
    highlight = (last,) if str(last).startswith("증감") else ()
    sc._style_sheet(ws, ncol, nrow, list(frame.columns), highlight, freeze_col=2)
    sc._fit_widths(ws, frame, ncol)

    fmt = sc.FMT_MN if unit != 1 else sc.FMT_COUNT
    center = Alignment(horizontal="center", vertical="center")
    left = Alignment(horizontal="left", vertical="center")
    right = Alignment(horizontal="right", vertical="center")
    for idx in range(1, ncol + 1):
        align = center if idx == 1 else (left if idx == 2 else right)
        for r in range(2, nrow + 2):
            cell = ws.cell(r, idx)
            if idx > 2:
                cell.number_format = fmt
            cell.alignment = align
    # 맨 아래 '총계' 행
    if nrow:
        for idx in range(1, ncol + 1):
            cell = ws.cell(nrow + 1, idx)
            cell.font = Font(bold=True, size=10)
            cell.fill = PatternFill("solid", fgColor="FFF3D6")


def add_line_chart(ws, summary, anchor_cell="G2"):
    """총계 추이 꺾은선. 매수/매도를 한 그림에 올린다."""
    from openpyxl.chart import LineChart, Reference

    ch = LineChart()
    ch.title = "전체 합계 추이 (USD mn)"
    ch.y_axis.title = "USD mn"
    ch.x_axis.title = "기준일"
    ch.height, ch.width = 9, 20
    nrow = len(summary)
    ch.add_data(Reference(ws, min_col=2, max_col=3, min_row=1, max_row=nrow + 1),
                titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=nrow + 1))
    for s in ch.series:
        s.smooth = False
    ws.add_chart(ch, anchor_cell)


def write_excel(rows, failed, st, ends, companies, log, tag):
    import pandas as pd

    OUT_DIR.mkdir(exist_ok=True)
    today = datetime.date.today().strftime("%Y%m%d")
    path = OUT_DIR / ("파생거래추이%s_%s.xlsx" % (tag, today))

    detail_cols = (["기준일", "사업자등록번호", "메모"]
                   + [c for c, _ in sc.COLMAP
                      if c not in ("_bizno", "사업자등록번호")])
    df = (pd.DataFrame(rows, columns=detail_cols) if rows
          else pd.DataFrame(columns=detail_cols))
    for c in df.columns:
        if c in sc.NUMERIC:
            df[c] = df[c].map(sc.to_num)
        if "최종갱신일" in c:
            # 자료가 없으면 '0000-00-00' 이 온다. 빈 칸이 읽기 낫다
            df[c] = df[c].replace({"0000-00-00": None, "00000000": None})
    df["결과코드"] = df["결과코드"].map(sc.norm_rspcd)

    names = _name_map(df, companies)
    order = [sc.fmt_bizno(b) for b, _ in companies]
    date_cols = [iso(d) for d in ends]

    frames = []
    for sheet, _title, srcs, unit in TREND_SHEETS:
        m = _pivot(df, companies, ends, srcs, unit)
        out = pd.DataFrame({"사업자등록번호": order,
                            "기업명": [names.get(b, "") for b in order]})
        for c in date_cols:
            out[c] = m[c].values
        if len(date_cols) >= 2:
            out["증감(%s→%s)" % (date_cols[0][5:], date_cols[-1][5:])] = (
                out[date_cols[-1]] - out[date_cols[0]])
        tot = {"사업자등록번호": "", "기업명": "총계"}
        for c in out.columns[2:]:
            tot[c] = out[c].sum(min_count=1)
        out = pd.concat([out, pd.DataFrame([tot])], ignore_index=True)
        frames.append((sheet, out, unit))

    # 상태 행렬: 추이표의 빈 칸이 '자료없음'인지 '조회실패'인지 구분해 준다
    stat = df.pivot_table(index="사업자등록번호", columns="기준일",
                          values="결과코드", aggfunc="last")
    stat = stat.reindex(index=order, columns=date_cols)
    status = pd.DataFrame({"사업자등록번호": order,
                           "기업명": [names.get(b, "") for b in order]})
    for c in date_cols:
        status[c] = [STATUS_TEXT.get(v, v if isinstance(v, str) else "미조회")
                     for v in stat[c].values]

    # 총계 요약 (차트용). 각 추이 시트의 마지막 '총계' 행을 가져다 쓴다.
    buy, sell, cnt = (f[1].iloc[-1] for f in frames)
    ok = (df[df["결과코드"].eq(sc.OK_CD)]
          .groupby("기준일")["사업자등록번호"].nunique())
    summary = pd.DataFrame({
        "기준일": date_cols,
        "매수 합계(USD mn)": [buy[c] for c in date_cols],
        "매도 합계(USD mn)": [sell[c] for c in date_cols],
        "총 건수": [cnt[c] for c in date_cols],
        "자료있는 기업수": [int(ok.get(c, 0)) for c in date_cols],
    })

    meta = pd.DataFrame([
        ("조회일시", datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("조회 시작일(고정)", iso(parse_ymd(st))),
        ("조회 종료일", "%s ~ %s (%d개 시점)"
         % (date_cols[0], date_cols[-1], len(date_cols))),
        ("대상 기업수", len(companies)),
        ("총 조회건수", len(rows)),
        ("조회 실패", len(failed)),
        ("", ""),
        ("읽는 법", "각 열은 %s 부터 그 날짜까지를 조회기간으로 준 결과입니다."
                    % iso(parse_ymd(st))),
        ("빈 칸", "해당 시점에 자료가 없거나(401) 조회에 실패한 경우입니다. "
                  "'상태' 시트에서 구분할 수 있습니다."),
        ("옵션 합산", "매수콜은 매수쪽, 매수풋은 매도쪽에만 더합니다. "
                      "매도콜·매도풋은 넣지 않습니다."),
    ], columns=["항목", "값"])

    fdf = (pd.DataFrame(failed, columns=["기준일", "사업자등록번호", "메모", "사유"])
           if failed else None)

    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        summary.to_excel(xw, sheet_name="요약", index=False)
        for sheet, out, _unit in frames:
            out.to_excel(xw, sheet_name=sheet, index=False)
        status.to_excel(xw, sheet_name="상태", index=False)
        df.to_excel(xw, sheet_name="상세", index=False)
        meta.to_excel(xw, sheet_name="조회정보", index=False)
        if fdf is not None:
            fdf.to_excel(xw, sheet_name="실패", index=False)

        ws = xw.sheets["요약"]
        sc.format_plain(ws, summary)
        for r in range(2, len(summary) + 2):
            ws.cell(r, 2).number_format = sc.FMT_MN
            ws.cell(r, 3).number_format = sc.FMT_MN
            ws.cell(r, 4).number_format = sc.FMT_COUNT
        add_line_chart(ws, summary)
        for sheet, out, unit in frames:
            _style_trend(xw.sheets[sheet], out, unit)
        sc.format_plain(xw.sheets["상태"], status)
        sc.format_grid(xw.sheets["상세"], df)
        sc.format_plain(xw.sheets["조회정보"], meta)
        if fdf is not None:
            sc.format_plain(xw.sheets["실패"], fdf)

    log("\n엑셀 저장 -> %s" % path)
    return path


# --------------------------------------------------------------------------
def main():
    global DELAY_SEC
    ap = argparse.ArgumentParser(
        description="조회 시작일을 올해 1월 1일로 고정하고, 종료일을 "
                    "2월부터 매월 옮겨가며 조회해 추이를 만든다.")
    ap.add_argument("--companies", default=None,
                    help="조회 대상 목록 CSV (기본: companies2.csv)")
    ap.add_argument("--year", type=int, default=None,
                    help="기준 연도 (기본: 올해)")
    ap.add_argument("--from", dest="st", default=None,
                    help="조회 시작일 YYYYMMDD (기본: 기준연도 1월 1일)")
    ap.add_argument("--anchor", choices=("first", "last"), default="first",
                    help="월별 종료일을 1일로 볼지 말일로 볼지 (기본: first)")
    ap.add_argument("--start-month", type=int, default=2,
                    help="첫 종료일의 월 (기본: 2 = 2월)")
    ap.add_argument("--include-today", action="store_true",
                    help="마지막 시점으로 오늘 날짜도 추가")
    ap.add_argument("--delay", type=float, default=DELAY_SEC,
                    help="요청 간격 초 (기본 2.5). 줄이면 서버가 막는다")
    ap.add_argument("--resume", action="store_true",
                    help="오늘 만들어진 중간저장 CSV 를 읽어 이어서 조회")
    a = ap.parse_args()

    DELAY_SEC = a.delay
    comp_path = pathlib.Path(a.companies) if a.companies else DEFAULT_COMPANIES
    if not comp_path.is_absolute() and not comp_path.exists():
        comp_path = BASE / comp_path
    tag = sc.out_tag_for(comp_path)

    year = a.year or datetime.date.today().year
    st = a.st or datetime.date(year, 1, 1).strftime("%Y%m%d")
    ends = period_ends(year, a.anchor, a.start_month,
                       include_today=a.include_today)

    LOG_DIR.mkdir(exist_ok=True)
    stamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    log = sc.Log(LOG_DIR / ("monthly_" + stamp + ".log"))

    log("목록 파일: %s" % comp_path.name)
    companies = sc.load_companies(comp_path, log)
    if not companies:
        log("%s 에 사업자등록번호가 없습니다." % comp_path)
        log.close()
        return 1
    if not ends:
        log("조회할 종료일이 없습니다. --year / --start-month 를 확인하세요.")
        log.close()
        return 1

    OUT_DIR.mkdir(exist_ok=True)
    part_path = OUT_DIR / ("_진행중_추이%s_%s.csv"
                           % (tag, datetime.date.today().strftime("%Y%m%d")))
    done, prev = load_done(part_path, log) if a.resume else (set(), [])

    left = len(companies) * len(ends) - len(done)
    mins = left * (DELAY_SEC + 0.7) / 60
    log("대상 %d개 기업 x %d개 시점 = %d회 조회 (예상 %.0f분)"
        % (len(companies), len(ends), left, mins))
    log("조회 시작일 (고정): %s" % iso(parse_ymd(st)))
    log("조회 종료일: %s" % ", ".join(iso(d) for d in ends))

    part_cols = (["기준일", "사업자등록번호", "메모"]
                 + [c for c, _ in sc.COLMAP
                    if c not in ("_bizno", "사업자등록번호")])
    append = a.resume and part_path.exists()
    part_fh = open(part_path, "a" if append else "w",
                   encoding="utf-8-sig", newline="")
    pw = csv.DictWriter(part_fh, fieldnames=part_cols, extrasaction="ignore")
    if part_fh.tell() == 0:
        pw.writeheader()

    state = {"user": None, "start": False, "rows": list(prev), "failed": [],
             "done": done, "writer": pw, "fh": part_fh}

    with sync_playwright() as p:
        ctx = p.chromium.launch_persistent_context(
            user_data_dir=str(sc.PROFILE_DIR), headless=False,
            locale="ko-KR", timezone_id="Asia/Seoul", viewport=None,
            args=["--disable-blink-features=AutomationControlled",
                  "--disable-popup-blocking", "--start-maximized"])
        ctx.add_init_script(sc.INJECT_JS)
        ctx.expose_binding("kfbStart",
                           lambda source: state.__setitem__("start", True))

        def on_response(resp):
            # 로그인 응답에서 사용자/기관 정보를 주워둔다. 조회 요청에 필요하다.
            if state["user"] or sc.SVC_LOGIN2 not in resp.url:
                return
            try:
                parsed = nx.parse_response(resp.body())
            except Exception:
                return
            for rws in parsed["datasets"].values():
                for r in rws:
                    if r.get("USER_ID") and r.get("USER_ORG_CD"):
                        state["user"] = {
                            "USER_ID": r["USER_ID"],
                            "USER_NM": r.get("USER_NM", ""),
                            "USER_ORG_CD": r["USER_ORG_CD"],
                            "USER_ORG_NM": r.get("USER_ORG_NM", ""),
                        }
                        log("로그인 확인: %s / %s (%s)"
                            % (r["USER_ID"], r.get("USER_NM"),
                               r.get("USER_ORG_NM")))
                        return

        ctx.on("response", on_response)
        page = ctx.pages[0] if ctx.pages else ctx.new_page()
        page.goto(sc.START_URL, wait_until="domcontentloaded")

        bar = "=" * 74
        log(bar)
        log(" 로그인하세요 (ID/PW + 추가 인증).")
        log(" 로그인이 끝나면 우측 하단 '자동 조회 시작' 버튼이 파랗게 켜집니다.")
        log(" 버튼을 누르면 %d개 기업을 %d개 시점으로 조회합니다. (약 %.0f분)"
            % (len(companies), len(ends), mins))
        log(" 중간에 끊기면 --resume 으로 이어받을 수 있습니다.")
        log(bar)

        def notify(text):
            for pg in ctx.pages:
                try:
                    pg.evaluate("t => window.__kfbMsg && window.__kfbMsg(t)", text)
                except Exception:
                    pass

        armed = False
        closed = {"v": False}
        ctx.on("close", lambda _: closed.__setitem__("v", True))
        ran = False
        started = time.time()
        while not closed["v"]:
            pages = [x for x in ctx.pages if not x.is_closed()]
            if not pages:
                break
            if state["user"] and not armed:
                armed = True
                txt = "준비 완료 - %d개 기업 x %d시점" % (len(companies), len(ends))
                for pg in pages:
                    try:
                        pg.evaluate(
                            "t => window.__kfbReady && window.__kfbReady(t)", txt)
                    except Exception:
                        pass
                log("버튼이 활성화됐습니다. 눌러주세요.")
            if state["start"] and state["user"]:
                log("\n" + bar)
                log(" 조회 시작")
                log(bar)
                ran = True
                try:
                    run_trend(ctx.request, state["user"], companies, st, ends,
                              log, notify, state)
                except KeyboardInterrupt:
                    log("\n중단됨. 지금까지 받은 결과만 저장합니다.")
                notify("완료: %d행 / 실패 %d"
                       % (len(state["rows"]), len(state["failed"])))
                break
            try:
                pages[0].wait_for_timeout(400)
            except Exception:
                break
            if time.time() - started > 6 * 3600:
                log("6시간이 지나 종료합니다.")
                break

        part_fh.close()
        if ran and state["rows"]:
            log("  (중간저장: %s)" % part_path.name)
            write_excel(state["rows"], state["failed"], st, ends, companies,
                        log, tag)
            log("총 %d행 / 실패 %d" % (len(state["rows"]), len(state["failed"])))
            for d, b, m, why in state["failed"]:
                log("  실패: %s %s %s - %s" % (d, b, m, why))
            try:
                pages = [x for x in ctx.pages if not x.is_closed()]
                if pages:
                    pages[0].wait_for_timeout(4000)
            except Exception:
                pass
        else:
            log("조회를 실행하지 않고 종료했습니다.")

        try:
            ctx.close()
        except Exception:
            pass

    log.close()
    return 0


if __name__ == "__main__":
    sys.exit(main())
