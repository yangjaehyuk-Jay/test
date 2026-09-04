# -*- coding: utf-8 -*-
"""
2단계: 파생상품 정보 일괄 조회

브라우저를 띄우면 사람이 로그인한다. 로그인이 끝나면 화면 우측 하단에
'자동 조회 시작' 버튼이 활성화된다. 누르면 companies.csv 의 사업자등록번호를
차례로 조회해서 엑셀로 저장한다.

조회는 화면을 클릭하지 않고, 화면이 쓰는 것과 같은 요청을 그대로 보낸다.
    POST /bizkiR0101SelectList.do   (파생상품 조회)
    POST /bizcmEventLoggerInsert.do (사이트 감사로그. 수동 조회와 동일하게 남긴다)

실행:  python scrape.py
"""
import sys, re, csv, json, time, datetime, pathlib, argparse
import unicodedata

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

from playwright.sync_api import sync_playwright
import nexacro as nx

BASE = pathlib.Path(__file__).parent
PROFILE_DIR = BASE / "browser_profile"
OUT_DIR = BASE / "output"
LOG_DIR = BASE / "logs"
COMPANIES = BASE / "companies.csv"

# 목록 파일마다 결과 파일을 따로 쓰기 위한 꼬리표.
# companies.csv -> "" (파생거래_20260901.xlsx)
# companies1.csv -> "1" (파생거래1_20260901.xlsx)
OUT_TAG = ""

HOST = "https://bizinfo.kfb.or.kr"
START_URL = HOST + "/biz/index.html"
POPUP_URL = (HOST + "/biz/popup.html"
             "?formname=BIZKI%3A%3ABIZKI_T0101.xfdl"
             "&framename=KI01_BIZKI_T0101.xfdl")
SVC_SELECT = HOST + "/bizkiR0101SelectList.do"
SVC_LOG = HOST + "/bizcmEventLoggerInsert.do"
SVC_LOGIN2 = "bizliS0102Login2.do"

DELAY_SEC = 2.5        # 요청 간 간격. 서버가 연속 호출을 차단하므로 필수
RETRY = 3
RETRY_WAIT = 8

XHR_HEADERS = {
    "content-type": "text/xml",
    "accept": "application/xml, text/xml, */*",
    "x-requested-with": "XMLHttpRequest",
    "cache-control": "no-cache, no-store",
    "pragma": "no-cache",
    "referer": POPUP_URL,
}

# 응답 컬럼 -> 엑셀 열 이름.
# FX=선물환, OP=옵션, CRS=통화스왑. BAL=잔액, VLD_CNT=유효건수.
COLMAP = [
    ("사업자등록번호", "_bizno"),
    ("기업명", "KIKO_CUST_NM"),
    ("결과코드", "RSP_CD"),
    ("결과내용", "RSP_NM"),
    ("선물환_건수", "FX_VLD_CNT"),
    ("선물환_매수잔액", "FX_BUY_BAL"),
    ("선물환_매도잔액", "FX_SELL_BAL"),
    ("선물환_최종갱신일", "FX_TRN_LST_UPD_DT"),
    ("옵션_건수", "OP_VLD_CNT"),
    ("옵션_매수잔액", "OP_BUY_BAL"),
    ("옵션_매도잔액", "OP_SELL_BAL"),
    ("옵션_매수콜", "OP_BUY_CALL_BAL"),
    ("옵션_매수풋", "OP_BUY_PUT_BAL"),
    ("옵션_매도콜", "OP_SELL_CALL_BAL"),
    ("옵션_매도풋", "OP_SELL_PUT_BAL"),
    ("옵션_최종갱신일", "OP_TRN_LST_UPD_DT"),
    ("통화스왑_건수", "CRS_VLD_CNT"),
    ("통화스왑_매수잔액", "CRS_BUY_BAL"),
    ("통화스왑_매도잔액", "CRS_SELL_BAL"),
    ("통화스왑_최종갱신일", "CRS_TRN_LST_UPD_DT"),
    ("USD환산금액", "BNS_AMT_USD_CVT_AMT"),
    ("조회일시", "RETRIEVE_DTTM"),
]
# 엑셀 맨 끝에 붙이는 합계 열.
# 옵션은 '매수' 포지션만 넣는다. 매수콜은 사실상 매수 포지션, 매수풋은 매도
# 포지션이므로 각각 매수/매도 쪽에 더한다. 매도콜·매도풋은 넣지 않는다.
# 원본 잔액은 USD 단위라 자릿수가 커서 읽기 어렵다. 합계 열만 백만 단위로 줄인다.
SUM_UNIT = 1_000_000
SUMCOLS = [
    ("선물환매수+통화스왑매수+옵션콜매수_잔액(USD mn)",
     ("선물환_매수잔액", "통화스왑_매수잔액", "옵션_매수콜")),
    ("선물환매도+스왑매도+옵션풋매수_잔액(USD mn)",
     ("선물환_매도잔액", "통화스왑_매도잔액", "옵션_매수풋")),
]
# 숫자 열은 이름이 아니라 응답 키로 고른다. 이름으로 "잔액"만 찾으면
# 옵션_매수콜(OP_BUY_CALL_BAL) 처럼 '잔액'이 안 붙은 열이 문자열로 남는다.
COUNTCOLS = {c for c, k in COLMAP if k.endswith("_CNT")}
NUMERIC = {c for c, k in COLMAP
           if k.endswith("_BAL") or k.endswith("_CNT") or k.endswith("_AMT")}
NUMERIC |= {c for c, _ in SUMCOLS}

# 조회 자체가 실패한 행의 결과코드. 서버가 주는 값(000/401)과 겹치지 않게 쓴다.
ERR_CD = "ERR"
OK_CD = "000"          # 정상. 이 값이 아니면 결과가 없는 행이다
# 결과가 없는 행의 합계 열에는 숫자 대신 이 문구를 넣는다
NO_RESULT_TEXT = "붙이지 마시오"


def norm_rspcd(v):
    """결과코드를 3자리 문자열로 되돌린다.

    엑셀을 거치면 '000' 이 숫자 0 으로 바뀐다. 그대로 두면 정상 행까지
    '결과 없음' 으로 잘못 잡히므로, 숫자로 온 값은 0 을 채워 되돌린다.
    """
    if v is None:
        return None
    s = str(v).strip()
    if s in ("", "nan", "None"):
        return None
    if s.replace(".0", "").isdigit():
        return s.replace(".0", "").zfill(3)
    return s

INJECT_JS = """
(() => {
  if (window.top !== window || window._popup) return;
  var add = function () {
    if (!document.body || document.getElementById('__kfb_run')) return;
    var box = document.createElement('div');
    box.id = '__kfb_run';
    box.style.cssText = 'position:fixed;right:14px;bottom:14px;z-index:2147483647;'
      + 'font:13px/1.5 sans-serif;text-align:right;';
    var msg = document.createElement('div');
    msg.id = '__kfb_msg';
    msg.style.cssText = 'margin-bottom:6px;padding:6px 10px;background:#222;'
      + 'color:#fff;border-radius:4px;opacity:.9;';
    msg.textContent = '\\ub85c\\uadf8\\uc778 \\ub300\\uae30 \\uc911...';
    var btn = document.createElement('button');
    btn.id = '__kfb_btn';
    btn.disabled = true;
    btn.textContent = '\\u25b6 \\uc790\\ub3d9 \\uc870\\ud68c \\uc2dc\\uc791';
    btn.style.cssText = 'padding:11px 16px;background:#999;color:#fff;border:0;'
      + 'border-radius:5px;cursor:not-allowed;font-size:14px;'
      + 'box-shadow:0 2px 8px rgba(0,0,0,.35);';
    btn.onclick = function () {
      btn.disabled = true; btn.style.background = '#999';
      btn.style.cursor = 'not-allowed';
      window.kfbStart();
    };
    box.appendChild(msg); box.appendChild(btn);
    document.body.appendChild(box);
  };
  if (document.readyState === 'loading')
    document.addEventListener('DOMContentLoaded', add);
  else add();
  setTimeout(add, 3000); setTimeout(add, 8000);

  window.__kfbReady = function (text) {
    var b = document.getElementById('__kfb_btn');
    var m = document.getElementById('__kfb_msg');
    if (m) m.textContent = text;
    if (b) { b.disabled = false; b.style.background = '#0796ec';
             b.style.cursor = 'pointer'; }
  };
  window.__kfbMsg = function (text) {
    var m = document.getElementById('__kfb_msg');
    if (m) m.textContent = text;
  };
})();
"""


class Log:
    def __init__(self, path):
        self.fh = open(path, "w", encoding="utf-8")
        self.path = path

    def __call__(self, msg=""):
        s = str(msg)
        self.fh.write(s + "\n")
        self.fh.flush()
        try:
            print(s)
        except Exception:
            print(s.encode("ascii", "replace").decode("ascii"))

    def close(self):
        self.fh.close()


# --------------------------------------------------------------------------
def read_text_any(path):
    """엑셀/메모장이 저장한 CP949 도, UTF-8 도 읽는다."""
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "cp949", "euc-kr"):
        try:
            return raw.decode(enc), enc
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace"), "utf-8(손상)"


def normalize_bizno(raw):
    """사업자등록번호 문자열을 10자리 숫자로 만든다.

    엑셀은 '0124011628' 을 숫자로 보고 앞의 0 을 떼어 '124011628' 로 저장한다.
    사업자등록번호는 언제나 10자리이므로, 짧게 들어온 값은 앞을 0 으로 채운다.
    ="0124011628" / '0124011628 / 012-40-11628 같은 표기도 모두 받는다.
    반환: (10자리, 보정메모) 또는 (None, 사유)
    """
    digits = re.sub(r"\D", "", raw)
    if not digits:
        return None, "숫자 없음"
    if len(digits) == 10:
        return digits, ""
    if len(digits) > 10:
        return None, "10자리 초과(%d자리)" % len(digits)
    if len(digits) >= 8:
        # 엑셀이 떼어낸 선행 0 을 되돌린다 (최대 2자리까지만)
        return digits.zfill(10), "선행 0 복원: %s -> %s" % (digits, digits.zfill(10))
    return None, "10자리 아님(%d자리)" % len(digits)


def bizno_checksum_ok(d):
    """국세청 사업자등록번호 검증코드(마지막 자리) 확인.

    앞 9자리에 가중치 1,3,7,1,3,7,1,3,5 를 곱해 더하고
    9번째 자리 x 5 의 십의 자리를 더한 뒤, 10 의 보수가 마지막 자리와 같아야 한다.
    오타/자릿수 누락을 조회 전에 걸러내는 용도다. 실패해도 조회는 막지 않는다.
    """
    if len(d) != 10 or not d.isdigit():
        return False
    w = (1, 3, 7, 1, 3, 7, 1, 3, 5)
    s = sum(int(d[i]) * w[i] for i in range(9)) + (int(d[8]) * 5) // 10
    return (10 - s % 10) % 10 == int(d[9])


def load_companies(path, log=print):
    """목록 CSV 를 읽는다.

    첫 열이 사업자등록번호, 둘째 열(선택)이 기업명, 셋째 열(선택)이 분류.
    둘째·셋째 열은 합쳐서 메모가 된다 -> "삼성중공업(주) (GC)".
    엑셀 `메모` 열과 실행 로그에 그대로 나온다.
    """
    if not path.exists():
        return []
    out = []
    text, enc = read_text_any(path)
    if enc != "utf-8-sig":
        log("  %s 인코딩: %s" % (path.name, enc))
    import io as _io
    if True:
        for lineno, row in enumerate(csv.reader(_io.StringIO(text)), 1):
            if not row:
                continue
            raw = row[0].strip()
            if not raw or raw.startswith("#"):
                continue
            digits, note = normalize_bizno(raw)
            if digits is None:
                log("  %d행 건너뜀 (%s): %s" % (lineno, note, raw))
                continue
            if note:
                log("  %d행 %s" % (lineno, note))
            if not bizno_checksum_ok(digits):
                log("  %d행 경고: 검증코드 불일치 %s - 번호가 틀렸을 수 있습니다"
                    % (lineno, fmt_bizno(digits)))
            name = row[1].strip() if len(row) > 1 else ""
            tag = row[2].strip() if len(row) > 2 else ""
            memo = "%s (%s)" % (name, tag) if name and tag else (name or tag)
            out.append((digits, memo))
    # 중복 제거 (순서 유지)
    seen, uniq = set(), []
    for d, m in out:
        if d not in seen:
            seen.add(d)
            uniq.append((d, m))
    return uniq


def fmt_bizno(d):
    return "%s-%s-%s" % (d[:3], d[3:5], d[5:])


def to_num(v):
    if v is None:
        return None
    s = str(v).replace(",", "").strip()
    if s in ("", "-"):
        return None
    try:
        return float(s)
    except ValueError:
        return None


def post_xml(req_ctx, url, body, log):
    """XML POST + 재시도. 서버가 연속 호출을 차단하므로 물러섰다 다시 친다."""
    last = None
    for attempt in range(1, RETRY + 1):
        try:
            resp = req_ctx.post(url, data=body, headers=XHR_HEADERS, timeout=60000)
            text = resp.text()
            if resp.status != 200:
                last = "HTTP %s" % resp.status
            elif "<Root" not in text:
                # 차단되면 '일시적인 오류' 안내 HTML 이 온다
                last = "XML 아님 (차단 추정, %d바이트)" % len(text)
            else:
                return text
        except Exception as e:
            last = str(e)[:120]
        if attempt < RETRY:
            log("      재시도 %d/%d (%s) - %d초 대기"
                % (attempt, RETRY - 1, last, RETRY_WAIT * attempt))
            time.sleep(RETRY_WAIT * attempt)
    raise RuntimeError(last or "알 수 없는 실패")


def audit_log(req_ctx, user, act, log):
    """사이트 감사로그를 수동 조회와 동일하게 남긴다."""
    body = nx.build_request(
        {"ds_log": [{
            "USER_ID": user["USER_ID"], "USER_NM": user["USER_NM"],
            "SCREEN_ID": "/BIZKI/BIZKI_R0101",
            "SCREEN_NM": "파생상품(TAB_01)", "ACT_NM": act,
        }]},
        params={"npPfsHost": "127.0.0.1", "npPfsPort": "14440"})
    try:
        post_xml(req_ctx, SVC_LOG, body, log)
    except Exception as e:
        log("      (감사로그 실패: %s)" % e)


def query_one(req_ctx, user, bizno, st, ed, log):
    body = nx.build_request(
        {"ds_search": [{
            "USER_ID": user["USER_ID"],
            "USER_NM": user["USER_NM"],
            "KIKO_ID_NO_TP_CD": "2",              # 2 = 사업자등록번호
            "KIKO_ID_NO": "999" + bizno,          # 화면이 만드는 조회요청 식별번호
            "INQ_DV_CD": "1",
            "INQ_ST_DT": st,
            "INQ_ED_DT": ed,
            "IS_IN_USER": "N",
            "USER_ORG_CD": user["USER_ORG_CD"],
            "USER_ORG_NM": user["USER_ORG_NM"],
        }]},
        params={"npPfsHost": "127.0.0.1", "npPfsPort": "14440"})
    text = post_xml(req_ctx, SVC_SELECT, body, log)
    parsed = nx.parse_response(text)
    nx.check(parsed, text)
    etc = (parsed["datasets"].get("ds_etc") or [{}])[0]
    lst = (parsed["datasets"].get("ds_list") or [{}])[0]
    merged = dict(lst)
    merged.update({k: v for k, v in etc.items() if v not in (None, "")})
    # 기업명이 전각문자로 온다 (ＳＡＭＹＡＮＧ　ＣＯＲＰ．) -> 반각으로 정리
    nm = merged.get("KIKO_CUST_NM")
    if nm:
        merged["KIKO_CUST_NM"] = unicodedata.normalize("NFKC", nm).strip()
    return merged


PARTIAL = {"rows": [], "failed": []}


def run_batch(req_ctx, user, companies, st, ed, log, notify):
    """한 건씩 조회한다. 중간에 끊겨도 결과를 잃지 않도록 즉시 CSV 에 적는다."""
    rows, failed = PARTIAL["rows"], PARTIAL["failed"]
    total = len(companies)
    OUT_DIR.mkdir(exist_ok=True)
    part_path = OUT_DIR / ("_진행중%s_%s.csv"
                           % (OUT_TAG, datetime.date.today().strftime("%Y%m%d")))
    cols = ["사업자등록번호", "메모"] + [c for c, _ in COLMAP if c != "_bizno"
                                        and c != "사업자등록번호"]
    part = open(part_path, "w", encoding="utf-8-sig", newline="")
    pw = csv.DictWriter(part, fieldnames=cols, extrasaction="ignore")
    pw.writeheader()

    audit_log(req_ctx, user, "화면오픈", log)
    for i, (bizno, memo) in enumerate(companies, 1):
        label = fmt_bizno(bizno) + (" (%s)" % memo if memo else "")
        log("  [%2d/%d] %s" % (i, total, label))
        notify("조회 중 %d/%d  %s" % (i, total, fmt_bizno(bizno)))
        try:
            data = query_one(req_ctx, user, bizno, st, ed, log)
            audit_log(req_ctx, user, "내역조회", log)
            rec = {"_bizno": fmt_bizno(bizno)}
            for col, key in COLMAP:
                if key == "_bizno":
                    continue
                rec[col] = data.get(key)
            rec["사업자등록번호"] = fmt_bizno(bizno)
            if memo:
                rec["메모"] = memo
            rows.append(rec)
            pw.writerow(rec); part.flush()
            log("       -> %s / %s | 선물환 %s건, 옵션 %s건, 스왑 %s건"
                % (data.get("RSP_CD"), (data.get("KIKO_CUST_NM") or "").strip(),
                   data.get("FX_VLD_CNT"), data.get("OP_VLD_CNT"),
                   data.get("CRS_VLD_CNT")))
        except Exception as e:
            log("       -> 실패: %s" % e)
            failed.append((fmt_bizno(bizno), memo, str(e)[:160]))
            # 실패해도 목록에서 빠지지 않게 빈 행을 남긴다. 원본 순서를 지켜야
            # 나중에 대조할 수 있고, 조회가 안 된 기업이 조용히 사라지지 않는다.
            rec = {"사업자등록번호": fmt_bizno(bizno),
                   "결과코드": ERR_CD, "결과내용": str(e)[:160]}
            if memo:
                rec["메모"] = memo
            rows.append(rec)
            pw.writerow(rec); part.flush()
        if i < total:
            time.sleep(DELAY_SEC)
    part.close()
    log("  (중간저장: %s)" % part_path.name)
    return rows, failed


# ---- 엑셀 서식 -------------------------------------------------------------
# 잔액/금액은 1의 자리까지. 소수점은 버리지 않고 반올림해서 보여만 준다
# (셀에는 원래 값이 그대로 있으므로 합계·검산에는 영향이 없다).
FMT_MONEY = '#,##0;[Red]-#,##0;"-"'
FMT_COUNT = '#,##0;-#,##0;"-"'
# 백만 단위 합계 열. 여기서는 소수점 두 자리가 있어야 자릿수가 살아난다
FMT_MN = '#,##0.00;[Red]-#,##0.00;"-"'

C_HEAD = "1F3864"        # 머리글 배경 (진한 남색)
C_HEAD_SUM = "1F6B3B"    # 합계 열 머리글 (진한 초록)
C_BAND = "F4F6FA"        # 짝수 행 배경
C_SUM = "EAF4EC"         # 합계 열 배경
C_LINE = "D0D7E5"        # 격자선

# 이 열부터 새 그룹이 시작된다. 왼쪽에 굵은 선을 넣어 묶음을 눈에 보이게 한다.
GROUP_START = {"선물환_건수", "옵션_건수", "통화스왑_건수", "USD환산금액"}


def _style_sheet(ws, ncol, nrow, headers, sumcols=(), freeze_col=1):
    """머리글 + 격자 + 줄무늬 + 정렬. 모든 시트에 공통으로 쓴다."""
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    thin = Side(style="thin", color=C_LINE)
    thick = Side(style="medium", color="9AA7BD")
    head_font = Font(bold=True, color="FFFFFF", size=10)
    band = PatternFill("solid", fgColor=C_BAND)
    sumfill = PatternFill("solid", fgColor=C_SUM)

    # 제일 긴 제목이 몇 줄로 접히는지 보고 머리글 높이를 잡는다
    longest = max([_disp_len(h) for h in headers] or [10])
    ws.row_dimensions[1].height = max(34, min(4, -(-longest // 22)) * 15 + 8)
    for c in range(1, ncol + 1):
        name = headers[c - 1]
        is_sum = name in sumcols
        h = ws.cell(1, c)
        h.font = head_font
        h.fill = PatternFill("solid",
                             fgColor=C_HEAD_SUM if is_sum else C_HEAD)
        h.alignment = Alignment(horizontal="center", vertical="center",
                                wrap_text=True)
        left = thick if (name in GROUP_START or is_sum) else thin
        h.border = Border(left=left, right=thin, top=thin, bottom=thin)
        for r in range(2, nrow + 2):
            cell = ws.cell(r, c)
            cell.border = Border(left=left, right=thin, top=thin, bottom=thin)
            if is_sum:
                cell.fill = sumfill
                cell.font = Font(bold=True, size=10)
            elif r % 2 == 1:
                cell.fill = band

    if nrow:
        ws.auto_filter.ref = "A1:%s%d" % (ws.cell(1, ncol).column_letter, nrow + 1)
    ws.freeze_panes = ws.cell(2, freeze_col + 1).coordinate


def _disp_len(s):
    """엑셀 열 너비용 글자수. 한글·한자는 두 칸을 먹는다."""
    return sum(2 if unicodedata.east_asian_width(ch) in "WF" else 1
               for ch in str(s))


def _fit_widths(ws, df, ncol):
    for idx in range(1, ncol + 1):
        name = str(df.columns[idx - 1])
        # 머리글은 줄바꿈되므로 가장 긴 낱말만 있으면 된다.
        # '_' 와 '+' 뒤에서도 끊어 읽는다 (합계 열 제목이 길다)
        words = name.replace("_", "_ ").replace("+", "+ ").split()
        head_w = max([_disp_len(w) for w in words] or [8])
        body = df.iloc[:, idx - 1].dropna().astype(str).head(80)
        body_w = max([_disp_len(v) for v in body] or [0])
        if name in NUMERIC:                       # 천단위 쉼표 자리
            body_w = min(body_w, 15)
        width = min(max(head_w + 3, body_w + 3, 9), 26)
        ws.column_dimensions[ws.cell(1, idx).column_letter].width = width


def split_dttm(df, src="조회일시", parts=("조회일자", "조회시각")):
    """'2026-08-31 15:40:35' 한 열을 날짜 / 시각 두 열로 쪼갠다.

    자리는 원래 열이 있던 그 자리. 문자열로 둔다 (ISO 형식이라 그대로 정렬된다).
    """
    if src not in df.columns:
        return df
    at = list(df.columns).index(src)
    s = df[src].astype(str).str.strip()
    s = s.mask(s.isin(("", "nan", "None", "NaT")))
    d = s.str.slice(0, 10)
    t = s.str.slice(11).str.strip().replace("", None)
    df = df.drop(columns=[src])
    df.insert(at, parts[0], d)
    df.insert(at + 1, parts[1], t)
    return df


def format_grid(ws, df, sumnames=()):
    """'파생거래' 시트 서식. 새로 만들 때도, 기존 파일을 다시 꾸밀 때도 쓴다."""
    from openpyxl.styles import Alignment, Font

    ncol, nrow = len(df.columns), len(df)
    # 기업명까지 고정해 두면 오른쪽으로 밀어도 어느 회사인지 보인다
    freeze = (list(df.columns).index("기업명") + 1
              if "기업명" in df.columns else 1)
    _style_sheet(ws, ncol, nrow, list(df.columns), sumnames, freeze)
    _fit_widths(ws, df, ncol)

    center = Alignment(horizontal="center", vertical="center")
    left = Alignment(horizontal="left", vertical="center")
    right = Alignment(horizontal="right", vertical="center")
    sumset = {n for n, _ in SUMCOLS}
    for idx, c in enumerate(df.columns, 1):
        if c in NUMERIC:
            if c in sumset:
                fmt = FMT_MN
            elif c in COUNTCOLS:
                fmt = FMT_COUNT
            else:
                fmt = FMT_MONEY
            align = right
        elif (c in ("사업자등록번호", "결과코드", "조회일자", "조회시각")
              or "최종갱신일" in c):
            fmt, align = None, center
        else:
            fmt, align = None, left
        for r in range(2, nrow + 2):
            cell = ws.cell(r, idx)
            if fmt and not isinstance(cell.value, str):
                cell.number_format = fmt
            # 합계 열에 들어간 '붙이지 마시오' 는 가운데로 (숫자만 오른쪽)
            cell.alignment = center if (c in sumset
                                        and isinstance(cell.value, str)) else align

    # 조회 실패/무자료 행은 흐리게 (결과코드 000 이 아닌 행).
    # 엑셀을 거치며 '000' 이 숫자 0 이 되기도 하므로 norm_rspcd 로 맞춰 본다.
    if "결과코드" in df.columns:
        rc = list(df.columns).index("결과코드") + 1
        for r in range(2, nrow + 2):
            code = norm_rspcd(ws.cell(r, rc).value)
            if code in (None, OK_CD):
                continue
            for idx in range(1, ncol + 1):
                cur = ws.cell(r, idx).font
                ws.cell(r, idx).font = Font(bold=cur.bold, size=10,
                                            color="9098A8")


def format_plain(ws, df):
    """'조회정보' / '실패' 처럼 단순한 시트 서식."""
    from openpyxl.styles import Alignment
    ncol = len(df.columns)
    _style_sheet(ws, ncol, len(df), list(df.columns))
    _fit_widths(ws, df, ncol)
    left = Alignment(horizontal="left", vertical="center")
    for r in range(2, len(df) + 2):
        for c in range(1, ncol + 1):
            ws.cell(r, c).alignment = left


def write_excel(rows, failed, st, ed, log):
    import pandas as pd
    OUT_DIR.mkdir(exist_ok=True)
    today = datetime.date.today().strftime("%Y%m%d")
    path = OUT_DIR / ("파생거래%s_%s.xlsx" % (OUT_TAG, today))

    cols = [c for c, _ in COLMAP if c != "_bizno"]
    cols = ["사업자등록번호"] + [c for c in cols if c != "사업자등록번호"]
    if any("메모" in r for r in rows):
        cols.insert(1, "메모")
    df = pd.DataFrame(rows, columns=cols) if rows else pd.DataFrame(columns=cols)
    for c in df.columns:
        if c in NUMERIC:
            df[c] = df[c].map(to_num)
        if "최종갱신일" in c:
            # 자료가 없으면 '0000-00-00' 이 온다. 빈 칸이 읽기 낫다
            df[c] = df[c].replace({"0000-00-00": None, "00000000": None})
    df = split_dttm(df)
    if "결과코드" in df.columns:
        df["결과코드"] = df["결과코드"].map(norm_rspcd)

    # 합계 열 (맨 끝). 한쪽만 값이 있으면 그 값, 양쪽 다 비면 빈 칸으로 둔다.
    # 값 자체를 백만으로 나눠 넣는다. 정렬·차트가 표시 단위와 어긋나지 않게.
    # 결과가 없는 행(401·ERR)은 숫자 대신 경고 문구를 넣는다. 0 으로 두면
    # 다른 데 옮겨 붙일 때 잔액이 정말 0 인 기업과 구별이 안 된다.
    for name, srcs in SUMCOLS:
        v = df[list(srcs)].sum(axis=1, min_count=1) / SUM_UNIT
        if "결과코드" in df.columns:
            v = v.astype(object).where(df["결과코드"].eq(OK_CD), NO_RESULT_TEXT)
        df[name] = v
    sumnames = [n for n, _ in SUMCOLS]

    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        df.to_excel(xw, sheet_name="파생거래", index=False)
        # rows 에 실패 행도 들어 있으므로 rows 가 곧 요청 건수다.
        n_401 = sum(1 for r in rows if str(r.get("결과코드") or "") == "401")
        meta = pd.DataFrame([
            ("조회일시", datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
            ("조회기간", "%s ~ %s" % (st, ed)),
            ("요청 기업수", len(rows)),
            ("조회 성공", len(rows) - len(failed)),
            ("  - 자료 있음", len(rows) - len(failed) - n_401),
            ("  - 해당자료 없음(401)", n_401),
            ("조회 실패", len(failed)),
        ], columns=["항목", "값"])
        meta.to_excel(xw, sheet_name="조회정보", index=False)
        fdf = (pd.DataFrame(failed, columns=["사업자등록번호", "메모", "사유"])
               if failed else None)
        if fdf is not None:
            fdf.to_excel(xw, sheet_name="실패", index=False)

        format_grid(xw.sheets["파생거래"], df, sumnames)
        for name, sheet in (("조회정보", meta), ("실패", fdf)):
            if sheet is not None:
                format_plain(xw.sheets[name], sheet)
    log("\n엑셀 저장 -> %s" % path)
    return path


# --------------------------------------------------------------------------
def out_tag_for(path):
    """companies1.csv -> '1'. 목록마다 결과 파일이 겹치지 않게 한다."""
    m = re.match(r"^companies(.*)$", path.stem, re.I)
    tag = m.group(1) if m else path.stem
    return re.sub(r"[^0-9A-Za-z가-힣_-]", "", tag)


def main():
    global DELAY_SEC, OUT_TAG
    ap = argparse.ArgumentParser()
    ap.add_argument("--from", dest="st", default=None,
                    help="조회 시작일 YYYYMMDD (기본: 올해 1월 1일)")
    ap.add_argument("--to", dest="ed", default=None,
                    help="조회 종료일 YYYYMMDD (기본: 오늘)")
    ap.add_argument("--delay", type=float, default=DELAY_SEC)
    ap.add_argument("--companies", default=None,
                    help="조회 대상 목록 CSV (기본: companies.csv). "
                         "companies1.csv 를 주면 결과는 파생거래1_YYYYMMDD.xlsx")
    a = ap.parse_args()

    DELAY_SEC = a.delay
    comp_path = pathlib.Path(a.companies) if a.companies else COMPANIES
    if not comp_path.is_absolute():
        comp_path = (BASE / comp_path) if not comp_path.exists() else comp_path
    OUT_TAG = out_tag_for(comp_path)
    today = datetime.date.today()
    st = a.st or today.replace(month=1, day=1).strftime("%Y%m%d")
    ed = a.ed or today.strftime("%Y%m%d")

    LOG_DIR.mkdir(exist_ok=True)
    stamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    log = Log(LOG_DIR / ("scrape_" + stamp + ".log"))

    log("목록 파일: %s" % comp_path.name)
    companies = load_companies(comp_path, log)
    if not companies:
        log("%s 에 사업자등록번호가 없습니다: %s" % (comp_path.name, comp_path))
        log("한 줄에 하나씩, 10자리 숫자로 넣어주세요. (하이픈은 있어도 됩니다)")
        log.close()
        return 1
    log("조회 대상 %d개 기업 / 조회기간 %s ~ %s / 간격 %.1f초"
        % (len(companies), st, ed, DELAY_SEC))

    state = {"user": None, "start": False}

    with sync_playwright() as p:
        ctx = p.chromium.launch_persistent_context(
            user_data_dir=str(PROFILE_DIR), headless=False,
            locale="ko-KR", timezone_id="Asia/Seoul", viewport=None,
            args=["--disable-blink-features=AutomationControlled",
                  "--disable-popup-blocking", "--start-maximized"])
        ctx.add_init_script(INJECT_JS)
        ctx.expose_binding("kfbStart",
                           lambda source: state.__setitem__("start", True))

        def on_response(resp):
            # 로그인 응답에서 사용자/기관 정보를 주워둔다. 조회 요청에 필요하다.
            if state["user"] or SVC_LOGIN2 not in resp.url:
                return
            try:
                parsed = nx.parse_response(resp.body())
            except Exception:
                return
            for rows in parsed["datasets"].values():
                for r in rows:
                    if r.get("USER_ID") and r.get("USER_ORG_CD"):
                        state["user"] = {
                            "USER_ID": r["USER_ID"], "USER_NM": r.get("USER_NM", ""),
                            "USER_ORG_CD": r["USER_ORG_CD"],
                            "USER_ORG_NM": r.get("USER_ORG_NM", ""),
                        }
                        log("로그인 확인: %s / %s (%s)"
                            % (r["USER_ID"], r.get("USER_NM"), r.get("USER_ORG_NM")))
                        return

        ctx.on("response", on_response)
        page = ctx.pages[0] if ctx.pages else ctx.new_page()
        page.goto(START_URL, wait_until="domcontentloaded")

        bar = "=" * 74
        log(bar)
        log(" 로그인하세요 (ID/PW + 추가 인증).")
        log(" 로그인이 끝나면 우측 하단 '자동 조회 시작' 버튼이 파랗게 켜집니다.")
        log(" 버튼을 누르면 %d개 기업을 차례로 조회합니다." % len(companies))
        log(bar)

        def notify(text):
            for pg in ctx.pages:
                try:
                    pg.evaluate("t => window.__kfbMsg && window.__kfbMsg(t)", text)
                except Exception:
                    pass

        armed = False
        closed = {"v": False}
        ctx.on("close", lambda _: closed.__setitem__("v", True))
        rows = failed = None
        started = time.time()
        while not closed["v"]:
            pages = [x for x in ctx.pages if not x.is_closed()]
            if not pages:
                break
            if state["user"] and not armed:
                armed = True
                txt = "준비 완료 - %d개 기업" % len(companies)
                for pg in pages:
                    try:
                        pg.evaluate("t => window.__kfbReady && window.__kfbReady(t)", txt)
                    except Exception:
                        pass
                log("버튼이 활성화됐습니다. 눌러주세요.")
            if state["start"] and state["user"]:
                log("\n" + bar)
                log(" 조회 시작")
                log(bar)
                try:
                    rows, failed = run_batch(ctx.request, state["user"], companies,
                                             st, ed, log, notify)
                except KeyboardInterrupt:
                    log("\n중단됨. 지금까지 받은 결과만 저장합니다.")
                    rows, failed = PARTIAL["rows"], PARTIAL["failed"]
                notify("완료: %d건 중 성공 %d / 실패 %d"
                       % (len(rows), len(rows) - len(failed), len(failed)))
                break
            try:
                pages[0].wait_for_timeout(400)
            except Exception:
                break
            if time.time() - started > 3600:
                log("1시간이 지나 종료합니다.")
                break

        if rows is not None:
            path = write_excel(rows, failed, st, ed, log)
            log("%d건 중 성공 %d / 실패 %d"
                % (len(rows), len(rows) - len(failed), len(failed)))
            for b, m, why in failed:
                log("  실패: %s %s - %s" % (b, m, why))
            notify("저장 완료: %s" % path.name)
            try:
                pages = [x for x in ctx.pages if not x.is_closed()]
                if pages:
                    pages[0].wait_for_timeout(4000)
            except Exception:
                pass
        else:
            log("조회를 실행하지 않고 종료했습니다.")

        try:
            ctx.close()
        except Exception:
            pass

    log.close()
    return 0


if __name__ == "__main__":
    sys.exit(main())
